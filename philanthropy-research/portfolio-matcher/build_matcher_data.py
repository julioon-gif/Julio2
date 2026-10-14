#!/usr/bin/env python3
"""Build data.json for the Portfolio Mission Matcher.

Usage:
  python3 build_matcher_data.py CAPITAL_FRAMEWORK.xlsx [STAGE2_MATCH.xlsx] -o OUT/data.json

Inputs
  CAPITAL_FRAMEWORK.xlsx  Julio's Capital Framework workbook (v4 layout: 'Capital Map v2',
                          'CSR Tiering', 'Opportunities', 'Plan Details' sheets).
  STAGE2_MATCH.xlsx       optional 'Account Match' workbook (Stage 2 offering picks).
  offerings.json, account-research.json   next to this script (public-source research).
  ../prospects.md         the weekly PNW scan tracker (funding programs and deadlines).

The workbooks hold internal plans and contacts, so data.json is written outside the repo
(pass -o to a scratch folder) and only published inside the private artifact.
"""
import json, re, sys, argparse, datetime
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
ALIASES = {"hpe": "cray inc"}  # CSR Tiering renamed Cray to HPE


def key(name):
    if not name:
        return ""
    base = re.split(r"[(/]", str(name))[0]
    words = [w for w in re.findall(r"[a-z0-9]+", base.lower()) if w not in ("the",)]
    k = " ".join(words[:2])
    return ALIASES.get(k, k)


def rows(ws):
    for r in ws.iter_rows(values_only=True):
        if any(c not in (None, "") for c in r):
            yield [("" if c is None else (str(c).strip() if not isinstance(c, (int, float)) else c)) for c in r]


def find_header(ws, first):
    data = list(rows(ws))
    for i, r in enumerate(data):
        if str(r[0]).strip() == first:
            return r, data[i + 1:]
    raise SystemExit(f"header '{first}' not found in {ws.title}")


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def pockets(text):
    t = str(text or "")
    out = []
    for p, pat in [("Grant", r"grant"), ("Commercial", r"commercial|sponsor|cause"), ("Payroll", r"payroll|wpg|employee"),
                   ("Principal", r"principal|owner"), ("Systems", r"system|platform"), ("Strategic", r"strategic|co-invest|advocacy")]:
        if re.search(pat, t, re.I):
            out.append(p)
    return out


def load_capital(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    hdr, body = find_header(wb["Capital Map v2"], "Account")
    accts, order = {}, []
    for r in body:
        name = str(r[0])
        if not name or name.startswith("──") or name in ("Live findings from the cross-tabulation (the evidentiary layer)",):
            continue
        if len(r) < 14 or r[2] == "" and r[3] == "":
            continue  # findings rows
        rec = dict(zip(["name", "sector", "tier", "capFn", "pocketRaw", "fyTier", "renewal", "predictability", "multiYear",
                        "owner", "anchorPath", "champRisk", "volatility", "capitalRead"], r[:14]))
        if rec["tier"] in ("—", "-"):
            rec["tier"] = "Route-out"
        accts[key(name)] = rec
        order.append(key(name))
    # stop at findings block: those rows have text in col B but empty classification; filtered above
    hdr, body = find_header(wb["CSR Tiering"], "Account")
    for r in body:
        k = key(r[0])
        if k in accts:
            a = accts[k]
            a.update(status=r[2], csrRead=r[3], capacity=r[4], causeFit=r[5], nextStep=r[6], confidenceRead=r[7])
            a.setdefault("csrName", r[0])
    if "Opportunities" in wb.sheetnames:
        hdr, body = find_header(wb["Opportunities"], "Account name")
        for r in body:
            k = key(r[0])
            if k in accts and (r[1] or r[6]):
                accts[k].setdefault("opps", []).append({
                    "name": r[1], "stage": r[2], "ask": num(r[3]),
                    "askDate": r[4][:10] if isinstance(r[4], str) else "", "forecast": num(r[5]), "matching": r[6]})
    if "Plan Details" in wb.sheetnames:
        ws = wb["Plan Details"]
        data = list(rows(ws))
        for r in data:
            k = key(r[0])
            if k in accts and len(r) > 2 and str(r[1]).strip() in ("1", "2", "3", "1.0", "2.0", "3.0"):
                accts[k]["wave"] = r[2]
                accts[k]["planStatus"] = r[3]
                if len(r) > 19:
                    accts[k]["nextDated"] = r[19]
    return accts, order


def load_stage2(path, accts):
    wb = openpyxl.load_workbook(path, data_only=True)
    hdr, body = find_header(wb["Account Match"], "Account")
    for r in body:
        k = key(r[0])
        if k in accts:
            accts[k]["stage2"] = dict(zip(["acct", "tier", "pocket", "primary", "why", "secondary", "confidence", "low", "high",
                                           "qtr", "mo", "firstStep", "caveats"], r[:13]))
            for f in ("low", "high"):
                accts[k]["stage2"][f] = num(accts[k]["stage2"][f])


def load_tracker():
    md = (HERE.parent / "prospects.md").read_text()
    out = {}
    for line in md.splitlines():
        if not line.startswith("| ") or line.startswith("| Company") or line.startswith("|---"):
            continue
        c = [x.strip() for x in line.strip().strip("|").split(" | ")]
        if len(c) < 9:
            continue
        out[c[0]] = dict(company=c[0], hq=c[1], program=c[2], tier=c[3], angle=c[4], deadline=c[5], status=c[6], checked=c[7], src=c[8])
    return out


# Verified or best-known dates for tracker deadlines (parsed by hand from the Next deadline text).
SCAN_DATES = {
    "KeyBank": ("2026-10-31", "V", "Community Sponsorship requests close end of October"),
    "Walmart": ("2026-11-30", "V", "Spark Good Local Grants Cycle 3 closes"),
    "Kroger (Fred Meyer, QFC)": ("2026-12-04", "V", "Kroger Co. Foundation quarterly grant deadline"),
    "PacifiCorp (Pacific Power Foundation)": ("2026-12-15", "V", "Pacific Power Foundation Safety/Wellness cycle"),
    "Nordstrom": ("2027-01-01", "U", "Community Grants window Jan 1-31 (third-party)"),
    "PSE (Puget Sound Energy)": ("2027-01-05", "U", "Powerful Partnerships 2027 likely early Jan"),
    "Les Schwab Tire Centers": ("2027-02-01", "V", "Tires for Purpose 2027 cycle opens"),
    "Bank of America": ("2027-02-01", "U", "2027 RFP windows likely Feb"),
    "Numerica Credit Union": ("2027-02-01", "U", "Cycles reopen early 2027 (no exact date)"),
    "BECU": ("2027-04-06", "U", "People Helping People nominations likely ~April"),
    "Blue Cross of Idaho": ("2027-05-01", "U", "Employee Community Fund likely ~May"),
    "Gesa Credit Union": ("2027-06-01", "U", "Local Heroes grants open ~Jun"),
    "CareOregon": ("2027-06-07", "U", "LOI window ~Jun"),
    "Moda Health": ("2027-09-01", "U", "Annual Sep 1 request deadline (inferred)"),
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("capital")
    ap.add_argument("stage2", nargs="?")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()

    accts, order = load_capital(a.capital)
    if a.stage2:
        load_stage2(a.stage2, accts)
    research = json.loads((HERE / "account-research.json").read_text())
    offerings = json.loads((HERE / "offerings.json").read_text())
    tracker = load_tracker()
    rkeys = {key(n): (n, v) for n, v in research["accounts"].items()}
    used_tracker = set()

    out_accts = []
    for k in order:
        acc = accts[k]
        rn, rv = rkeys.get(k, (None, {}))
        acc["id"] = re.sub(r"[^a-z0-9]+", "-", k).strip("-")
        acc["pockets"] = pockets(acc.get("pocketRaw")) + ([p for p in pockets(acc.get("stage2", {}).get("pocket")) if p not in pockets(acc.get("pocketRaw"))] if acc.get("stage2") else [])
        acc["region"] = rv.get("region", "")
        acc["tags"] = rv.get("tags", [])
        acc["notes"] = rv.get("notes", [])
        acc["funding"] = list(rv.get("funding", []))
        t = tracker.get(rv.get("tracker", ""))
        if t:
            used_tracker.add(t["company"])
            acc["tracker"] = t
            d = SCAN_DATES.get(t["company"])
            if d and not any(f.get("date") == d[0] for f in acc["funding"]):
                acc["funding"].append({"what": d[2], "date": d[0], "kind": "funder", "mark": d[1], "src": t["src"]})
        if isinstance(acc.get("tier"), float):
            acc["tier"] = int(acc["tier"])
        out_accts.append(acc)

    scan = []
    for name, t in tracker.items():
        if name in used_tracker:
            continue
        d = SCAN_DATES.get(name)
        scan.append(dict(t, date=d[0] if d else "", dateMark=d[1] if d else "", dateWhat=d[2] if d else ""))

    data = {
        "generated": datetime.date.today().isoformat(),
        "sources": {"capital": Path(a.capital).name, "stage2": Path(a.stage2).name if a.stage2 else None,
                    "tracker": "philanthropy-research/prospects.md"},
        "units": offerings["impact_units"], "offerings": offerings["offerings"],
        "accounts": out_accts, "scan": scan, "acsEvents": research.get("acs_events", []),
    }
    Path(a.out).write_text(json.dumps(data, indent=1, default=str))
    print(f"accounts {len(out_accts)} · offerings {len(data['offerings'])} · scan prospects {len(scan)} · tracker matched {len(used_tracker)}")
    missing = [n for n in research["accounts"] if key(n) not in accts]
    if missing:
        print("research keys with no Capital Map row:", missing)
    nostage = [a["name"] for a in out_accts if a.get("tier") in (1, 2) and not a.get("stage2")]
    if nostage:
        print("Tier 1-2 without a Stage 2 row:", nostage)


if __name__ == "__main__":
    main()
